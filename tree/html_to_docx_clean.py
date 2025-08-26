# html_to_docx_clean.py
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
import os, re

# --------- tiny style utilities ---------

def parse_style(style_str: str) -> dict:
    """
    Convert 'font-size:12pt; color:#333; font-family: Times New Roman' into a dict.
    Keys/values are lowercased & stripped.
    """
    out = {}
    if not style_str:
        return out
    for part in style_str.split(";"):
        if ":" in part:
            k, v = part.split(":", 1)
            out[k.strip().lower()] = v.strip().lower()
    return out

def color_to_rgb(color_str: str) -> RGBColor | None:
    """Accepts '#rrggbb', 'rrggbb', or 'rgb(r,g,b)'. Returns RGBColor or None."""
    if not color_str:
        return None
    s = color_str.strip().lower()
    # rgb(12, 34, 56)
    m = re.match(r"rgb\(\s*(\d{1,3})\s*,\s*(\d{1,3})\s*,\s*(\d{1,3})\s*\)", s)
    if m:
        r, g, b = [max(0, min(255, int(x))) for x in m.groups()]
        return RGBColor(r, g, b)
    # #aabbcc or aabbcc
    s = s[1:] if s.startswith("#") else s
    if re.fullmatch(r"[0-9a-f]{6}", s):
        return RGBColor.from_string(s)
    return None

def apply_run_style(run, styles: dict):
    """Apply inline CSS-ish styles to a python-docx run."""
    # font-family
    fam = styles.get("font-family")
    if fam:
        # strip quotes and take the first family if multiple
        fam = fam.split(",")[0].strip().strip("'\"")
        run.font.name = fam

    # font-size
    size = styles.get("font-size")
    if size:
        # support: "12pt", "14 px" (rough conversion), "12"
        size = size.replace(" ", "")
        try:
            if size.endswith("pt"):
                run.font.size = Pt(float(size[:-2]))
            elif size.endswith("px"):
                # rough px->pt (96dpi): pt = px * 72/96 = px * 0.75
                run.font.size = Pt(float(size[:-2]) * 0.75)
            else:
                run.font.size = Pt(float(size))
        except ValueError:
            pass

    # color
    col = styles.get("color")
    rgb = color_to_rgb(col) if col else None
    if rgb:
        run.font.color.rgb = rgb

    # bold / italic / underline
    if styles.get("font-weight") in ("bold", "700", "800", "900"):
        run.bold = True
    if styles.get("font-style") == "italic":
        run.italic = True
    # text-decoration can have multiple, check for 'underline'
    deco = styles.get("text-decoration", "")
    if "underline" in deco:
        run.underline = True

def merge_styles(base: dict, override: dict) -> dict:
    """Shallow merge where override wins."""
    out = dict(base or {})
    out.update(override or {})
    return out

# --------- core conversion ---------

def html_to_docx(html_string: str, output_path: str):
    """
    Convert a subset of HTML to a .docx file, preserving alignment + basic inline styles.
    Handles: <p>, <span>, <strong>/<b>, <em>/<i>, <br>, and CSS for font/color/bold/italic/underline.
    """
    doc = Document()
    soup = BeautifulSoup(html_string, "html.parser")

    def render_node(node, para, inherited_styles: dict):
        """Recursively render node contents into 'para', applying inherited + node styles."""
        if isinstance(node, NavigableString):
            text = str(node)
            if text:
                run = para.add_run(text)
                apply_run_style(run, inherited_styles)
            return

        if not isinstance(node, Tag):
            return

        # local style map from this tag's inline style
        local_styles = parse_style(node.get("style", ""))

        # tag-implied styles
        tag = node.name.lower()
        if tag in ("strong", "b"):
            local_styles = merge_styles(local_styles, {"font-weight": "bold"})
        elif tag in ("em", "i"):
            local_styles = merge_styles(local_styles, {"font-style": "italic"})

        # merged styles to pass down to children
        merged = merge_styles(inherited_styles, local_styles)

        if tag == "br":
            # line break inside the same paragraph
            run = para.add_run()
            run.add_break()  # soft line break
            return

        # Recurse into children so we don't double-write text
        for child in node.children:
            render_node(child, para, merged)

    # map common alignment sources
    ALIGN_MAP = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }

    for p in soup.find_all("p"):
        para = doc.add_paragraph("")
        p_styles = parse_style(p.get("style", ""))

        # alignment via attribute or CSS
        align_attr = (p.get("align") or "").strip().lower()
        text_align_css = p_styles.get("text-align", "").strip().lower()
        if align_attr in ALIGN_MAP:
            para.alignment = ALIGN_MAP[align_attr]
        elif text_align_css in ALIGN_MAP:
            para.alignment = ALIGN_MAP[text_align_css]

        # Render only children (prevents duplicate text from descendants)
        for child in p.children:
            render_node(child, para, p_styles)

    # ensure directory exists
    out_dir = os.path.dirname(os.path.abspath(output_path))
    if out_dir and not os.path.exists(out_dir):
        os.makedirs(out_dir, exist_ok=True)

    doc.save(output_path)


# --------- example ----------
if __name__ == "__main__":
    # === Test case 1: basic sample ===
    html_string = """
    <p style="text-align: justify; font-family: 'Times New Roman'; font-size:12pt;">
      <span style="color:#000000">Pursuant to the Policy, the <strong>Insurers</strong> 
      do not assume any duty to defend or investigate.</span>
      <br>
      <span style="font-style: italic; text-decoration: underline;">This line is italic and underlined.</span>
      <span style="color: rgb(120, 40, 140);"> And this one is purple via rgb().</span>
    </p>
    """
    html_to_docx(html_string, "./output/output_basic.docx")

    # === Test case 2: Word-exported HTML from your screenshot ===
    html_string_long = """<p class="MsoNormal" style="margin:0in 0in 0.25pt; 
    text-align: justify; text-indent: -0.5pt; line-height: 16.48px;
    font-size: 12pt; font-family: Cambria, serif; color: #000000;">
    <span lang="EN" style="font-family: 'Times New Roman', serif;">
    Pursuant to (9A) of the Policy, the <strong>Insureds</strong> shall defend
    and contest any <strong>Claim</strong> made against them.&nbsp; 
    AIG does not assume any duty to defend or investigate.&nbsp; 
    AIG has the right, but not the obligation, to fully and effectively
    associate with and every <strong>Organization</strong> and 
    <strong>Insured Person</strong> in the defense and prosecution of any
    <strong>Claim</strong> or <strong>Pre-Claim Inquiry</strong>. 
    [...]</span></p>"""
    html_to_docx(html_string_long, "./output/output_long.docx")

    # === Test case 3: color and font styling ===
    html_string_colors = """
    <p style="color: #008000;">Green paragraph. 
      <span style="color: #FFA500;">Orange span</span> 
      <span style="color: #800080;">Purple span</span> 
      <span style="color: #000000;">Black span</span>
    </p>
    <p style="color: #FF00FF;">Magenta paragraph. 
      <span style="color: #00FFFF;">Cyan span</span> 
      <strong style="color: #FF0000;">Red bold</strong> 
      <em style="color: #0000FF;">Blue italic</em>
    </p>
    """
    html_to_docx(html_string_colors, "./output/output_colors.docx")

